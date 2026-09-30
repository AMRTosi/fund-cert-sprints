from __future__ import annotations

import json

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from sprint_cert_automation.services.cost_hours_export_service import CostHoursExportService


def test_export_cost_hours_table_to_json(tmp_path) -> None:
    workbook_path = tmp_path / "forecast.xlsx"
    output_root = tmp_path / "out"
    _build_workbook_with_table(
        workbook_path,
        table_name="TableHorasCosteQuincenas202609",
        headers=["reviewee", "horasQuincena1", "horasQuincena2"],
        rows=[
            [1001, 80, 72],
            ["1002", 79.0, 70.5],
            [None, None, None],
        ],
    )

    result = CostHoursExportService().run(
        workbook_path=workbook_path,
        table_name="TableHorasCosteQuincenas202609",
        output_file_name="2026-09",
        output_root=output_root,
    )

    assert result.rows_exported == 2
    assert result.output_path == output_root / "2026-09.json"

    payload = json.loads(result.output_path.read_text(encoding="utf-8"))
    assert payload == [
        {
            "reviewee": "1001",
            "horasQuincena1": "80",
            "horasQuincena2": "72",
        },
        {
            "reviewee": "1002",
            "horasQuincena1": "79",
            "horasQuincena2": "70.5",
        },
    ]


def test_export_fails_when_table_is_missing(tmp_path) -> None:
    workbook_path = tmp_path / "forecast.xlsx"
    output_root = tmp_path / "out"
    _build_workbook_with_table(
        workbook_path,
        table_name="TableHorasCosteQuincenas202609",
        headers=["reviewee", "horasQuincena1", "horasQuincena2"],
        rows=[["1001", 80, 72]],
    )

    with pytest.raises(ValueError, match="Table not found"):
        CostHoursExportService().run(
            workbook_path=workbook_path,
            table_name="TableHorasCosteQuincenas202610",
            output_file_name="2026-10",
            output_root=output_root,
        )


def test_export_fails_when_required_columns_are_missing(tmp_path) -> None:
    workbook_path = tmp_path / "forecast.xlsx"
    output_root = tmp_path / "out"
    _build_workbook_with_table(
        workbook_path,
        table_name="TableHorasCosteQuincenas202609",
        headers=["eid", "horasQuincena1", "horasQuincena2"],
        rows=[["1001", 80, 72]],
    )

    with pytest.raises(ValueError, match="Required columns missing"):
        CostHoursExportService().run(
            workbook_path=workbook_path,
            table_name="TableHorasCosteQuincenas202609",
            output_file_name="2026-09",
            output_root=output_root,
        )


def test_export_fails_when_table_name_format_is_invalid(tmp_path) -> None:
    workbook_path = tmp_path / "forecast.xlsx"
    output_root = tmp_path / "out"
    _build_workbook_with_table(
        workbook_path,
        table_name="TableHorasCosteQuincenas202609",
        headers=["reviewee", "horasQuincena1", "horasQuincena2"],
        rows=[["1001", 80, 72]],
    )

    with pytest.raises(ValueError, match="TableHorasCosteQuincenasAAAAMM"):
        CostHoursExportService().run(
            workbook_path=workbook_path,
            table_name="TablaCustom",
            output_file_name="2026-09",
            output_root=output_root,
        )


def _build_workbook_with_table(workbook_path, table_name: str, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "FY27_oct"

    start_row = 1
    start_col = 1

    for col_offset, header in enumerate(headers):
        worksheet.cell(row=start_row, column=start_col + col_offset, value=header)

    for row_offset, row_values in enumerate(rows, start=1):
        for col_offset, value in enumerate(row_values):
            worksheet.cell(row=start_row + row_offset, column=start_col + col_offset, value=value)

    end_row = start_row + max(len(rows), 1)
    end_col = start_col + len(headers) - 1
    ref = f"A1:{chr(ord('A') + end_col - 1)}{end_row}"

    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    worksheet.add_table(table)

    workbook.save(workbook_path)
    workbook.close()
