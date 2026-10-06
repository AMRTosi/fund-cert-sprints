from __future__ import annotations

import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Color, Font

from sprint_cert_automation.services.beeline_timesheet_validation_service import (
    APPROVED_FONT_RGB,
    BeelineTimesheetValidationService,
)


def test_validate_beeline_timesheets_end_to_end(tmp_path) -> None:
    workbook_path = tmp_path / "control.xlsx"
    input_json_path = tmp_path / "input.json"
    output_root = tmp_path / "out"

    _build_control_workbook(workbook_path)
    input_payload = [
        {
            "beelineRevieweeName": "de Manuel Ruiz, Alberto",
            "timePeriod": "9/28/2026-10/4/2026",
            "totalEstimatedAmount": "1.000,00",
        },
        {
            "beelineRevieweeName": "Perez Gomez, Ana",
            "timePeriod": "9/28/2026-10/4/2026",
            "totalEstimatedAmount": "1200",
        },
        {
            "beelineRevieweeName": "Sanchez Ruiz, Marta",
            "timePeriod": "9/28/2026-10/4/2026",
            "totalEstimatedAmount": "999",
        },
        {
            "beelineRevieweeName": "No Existe, Persona",
            "timePeriod": "9/28/2026-10/4/2026",
            "totalEstimatedAmount": "500",
        },
        {
            "beelineRevieweeName": "Lopez Martin, Raul",
            "timePeriod": "9/28/2026-10/4/2026",
            "totalEstimatedAmount": "300",
        },
        {
            "beelineRevieweeName": "Moreno Diaz, Elsa",
            "timePeriod": "9/28/2026-10/4/2026",
        },
    ]
    input_json_path.write_text(json.dumps(input_payload, ensure_ascii=False), encoding="utf-8")

    result = BeelineTimesheetValidationService().run(
        workbook_path=workbook_path,
        input_json_path=input_json_path,
        output_file_name="validation_2026_10",
        output_root=output_root,
    )

    assert result.validated_rows == 6
    assert result.approved_rows == 1
    assert result.rejected_rows == 3
    assert result.not_found_rows == 2

    output_payload = json.loads(result.output_path.read_text(encoding="utf-8"))
    assert output_payload == [
        {
            "revieweeBeelineName": "de Manuel Ruiz, Alberto",
            "timePeriod": "9/28/2026-10/4/2026",
            "validationStatus": "OK",
            "message": "",
            "beelineAction": "Approve",
        },
        {
            "revieweeBeelineName": "Perez Gomez, Ana",
            "timePeriod": "9/28/2026-10/4/2026",
            "validationStatus": "Failed",
            "message": "Timecard enviada por duplicado.",
            "beelineAction": "Reject",
        },
        {
            "revieweeBeelineName": "Sanchez Ruiz, Marta",
            "timePeriod": "9/28/2026-10/4/2026",
            "validationStatus": "Failed",
            "message": "Timecard incorrecta. Revisar y Rechzarla",
            "beelineAction": "Reject",
        },
        {
            "revieweeBeelineName": "No Existe, Persona",
            "timePeriod": "9/28/2026-10/4/2026",
            "validationStatus": "NotFound",
            "message": "Reviewee not Found.",
            "beelineAction": "NA",
        },
        {
            "revieweeBeelineName": "Lopez Martin, Raul",
            "timePeriod": "9/28/2026-10/4/2026",
            "validationStatus": "NotFound",
            "message": "totalEstimatedAmount no disponible en Excel. Revisar cálculo en el fichero.",
            "beelineAction": "NA",
        },
        {
            "revieweeBeelineName": "Moreno Diaz, Elsa",
            "timePeriod": "9/28/2026-10/4/2026",
            "validationStatus": "Failed",
            "message": "Input item missing required fields or empty values.",
            "beelineAction": "Reject",
        },
    ]

    workbook = load_workbook(workbook_path)
    try:
        worksheet = workbook["Timecards Aprobadas"]

        approved_cell = worksheet["B4"]
        assert approved_cell.font.bold is True
        assert approved_cell.font.color is not None
        assert approved_cell.font.color.type == "rgb"
        assert approved_cell.font.color.rgb == APPROVED_FONT_RGB

        duplicate_cell = worksheet["B5"]
        assert duplicate_cell.font.bold is True
        assert duplicate_cell.font.color is not None
        assert duplicate_cell.font.color.type == "rgb"
        assert duplicate_cell.font.color.rgb == APPROVED_FONT_RGB

        mismatch_cell = worksheet["B6"]
        assert mismatch_cell.font.bold is False
    finally:
        workbook.close()


def test_validate_beeline_timesheets_requires_array_input(tmp_path) -> None:
    workbook_path = tmp_path / "control.xlsx"
    input_json_path = tmp_path / "input.json"
    _build_control_workbook(workbook_path)
    input_json_path.write_text(json.dumps({"foo": "bar"}), encoding="utf-8")

    with pytest.raises(ValueError, match="input JSON must be an array"):
        BeelineTimesheetValidationService().run(
            workbook_path=workbook_path,
            input_json_path=input_json_path,
            output_file_name="validation_2026_10",
            output_root=tmp_path / "out",
        )


def test_validate_beeline_timesheets_numeric_comparison_works_with_us_and_eu_formats(tmp_path) -> None:
    workbook_path = tmp_path / "control.xlsx"
    input_json_path = tmp_path / "input.json"
    output_root = tmp_path / "out"

    _build_control_workbook(workbook_path)
    input_payload = [
        {
            "beelineRevieweeName": "de Manuel Ruiz, Alberto",
            "timePeriod": "9/28/2026-10/4/2026",
            "totalEstimatedAmount": "1000.00",
        },
        {
            "beelineRevieweeName": "Sanchez Ruiz, Marta",
            "timePeriod": "9/28/2026-10/4/2026",
            "totalEstimatedAmount": "1,200.00",
        },
    ]
    input_json_path.write_text(json.dumps(input_payload, ensure_ascii=False), encoding="utf-8")

    result = BeelineTimesheetValidationService().run(
        workbook_path=workbook_path,
        input_json_path=input_json_path,
        output_file_name="validation_2026_10",
        output_root=output_root,
    )

    assert result.validated_rows == 2
    assert result.approved_rows == 2

    output_payload = json.loads(result.output_path.read_text(encoding="utf-8"))
    assert output_payload[0]["validationStatus"] == "OK"
    assert output_payload[1]["validationStatus"] == "OK"


def test_validate_beeline_timesheets_matches_formula_header_with_leading_zero_dates(tmp_path) -> None:
    workbook_path = tmp_path / "control.xlsx"
    input_json_path = tmp_path / "input.json"
    output_root = tmp_path / "out"

    _build_control_workbook(workbook_path)
    input_payload = [
        {
            "beelineRevieweeName": "Sanchez Ruiz, Marta",
            "timePeriod": "10/05/2026–10/11/2026",
            "totalEstimatedAmount": "1275.00",
        }
    ]
    input_json_path.write_text(json.dumps(input_payload, ensure_ascii=False), encoding="utf-8")

    result = BeelineTimesheetValidationService().run(
        workbook_path=workbook_path,
        input_json_path=input_json_path,
        output_file_name="validation_2026_10",
        output_root=output_root,
    )

    assert result.validated_rows == 1
    assert result.approved_rows == 1

    output_payload = json.loads(result.output_path.read_text(encoding="utf-8"))
    assert output_payload[0]["validationStatus"] == "OK"

    workbook = load_workbook(workbook_path)
    try:
        worksheet = workbook["Timecards Aprobadas"]
        approved_cell = worksheet["C6"]
        assert approved_cell.font.bold is True
        assert approved_cell.font.color is not None
        assert approved_cell.font.color.type == "rgb"
        assert approved_cell.font.color.rgb == APPROVED_FONT_RGB
    finally:
        workbook.close()


def _build_control_workbook(workbook_path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Timecards Aprobadas"

    worksheet["A3"] = "reviewee"
    worksheet["B3"] = "9/28/2026-10/4/2026"
    worksheet["C3"] = '=TEXT(DATE(VALUE(RIGHT(LEFT(B3,FIND("-",B3)-1),4)),VALUE(LEFT(B3,FIND("/",B3)-1)),VALUE(MID(B3,FIND("/",B3)+1,FIND("/",B3,FIND("/",B3)+1)-FIND("/",B3)-1)))+7,"m/d/yyyy")&"-"&TEXT(DATE(VALUE(RIGHT(LEFT(B3,FIND("-",B3)-1),4)),VALUE(LEFT(B3,FIND("/",B3)-1)),VALUE(MID(B3,FIND("/",B3)+1,FIND("/",B3,FIND("/",B3)+1)-FIND("/",B3)-1)))+13,"m/d/yyyy")'

    worksheet["A4"] = "de Manuel Ruiz, Alberto"
    worksheet["B4"] = "1000"
    worksheet["C4"] = "1000"

    worksheet["A5"] = "Perez Gomez, Ana"
    worksheet["B5"] = "1200"
    worksheet["C5"] = "1200"
    worksheet["B5"].font = Font(bold=True, color=Color(rgb=APPROVED_FONT_RGB))

    worksheet["A6"] = "Sanchez Ruiz, Marta"
    worksheet["B6"] = "1200"
    worksheet["C6"] = "1275"

    worksheet["A7"] = "Lopez Martin, Raul"
    worksheet["B7"] = None
    worksheet["C7"] = None

    workbook.save(workbook_path)
    workbook.close()
