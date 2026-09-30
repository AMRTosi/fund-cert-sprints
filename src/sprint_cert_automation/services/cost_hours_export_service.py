from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import json
import re

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

TABLE_NAME_PATTERN = re.compile(r"^TableHorasCosteQuincenas\d{6}$")
REQUIRED_HEADERS = ("reviewee", "email", "subco", "horasQuincena1", "horasQuincena2")
REQUIRED_HEADER_KEYS = {header.casefold(): header for header in REQUIRED_HEADERS}


@dataclass(frozen=True)
class CostHoursExportResult:
    workbook_path: Path
    output_path: Path
    table_name: str
    rows_exported: int


class CostHoursExportService:
    """Export current period cost hours table rows to a JSON array file."""

    def run(
        self,
        workbook_path: Path,
        table_name: str,
        output_file_name: str,
        output_root: Path,
    ) -> CostHoursExportResult:
        self._validate_table_name(table_name)
        normalized_output_name = self._normalize_output_name(output_file_name)

        workbook = load_workbook(workbook_path, data_only=True)
        try:
            worksheet, table_ref = self._resolve_table(workbook, table_name)
            rows = self._read_rows(worksheet, table_ref)
        finally:
            workbook.close()

        output_root.mkdir(parents=True, exist_ok=True)
        output_path = output_root / f"{normalized_output_name}.json"
        output_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")

        return CostHoursExportResult(
            workbook_path=workbook_path,
            output_path=output_path,
            table_name=table_name,
            rows_exported=len(rows),
        )

    def _validate_table_name(self, table_name: str) -> None:
        if not TABLE_NAME_PATTERN.match(table_name):
            raise ValueError(
                "table-name must match pattern 'TableHorasCosteQuincenasAAAAMM' "
                "(example: TableHorasCosteQuincenas202609)"
            )

    def _normalize_output_name(self, output_file_name: str) -> str:
        name = output_file_name.strip()
        if name.lower().endswith(".json"):
            name = name[:-5]
        if not name:
            raise ValueError("output-file-name must not be empty")
        if any(char in name for char in "\\/:*?\"<>|"):
            raise ValueError("output-file-name contains invalid filesystem characters")
        return name

    def _resolve_table(self, workbook, table_name: str):
        for worksheet in workbook.worksheets:
            if table_name in worksheet.tables:
                table = worksheet.tables[table_name]
                return worksheet, table.ref
        raise ValueError(f"Table not found in workbook: {table_name}")

    def _read_rows(self, worksheet, table_ref: str) -> list[dict[str, str]]:
        min_col, min_row, max_col, max_row = range_boundaries(table_ref)
        headers = [worksheet.cell(row=min_row, column=col).value for col in range(min_col, max_col + 1)]

        header_map = self._build_header_map(headers, min_col)
        missing = [
            REQUIRED_HEADER_KEYS[header_key]
            for header_key in REQUIRED_HEADER_KEYS
            if header_key not in header_map
        ]
        if missing:
            raise ValueError(f"Required columns missing in table: {', '.join(missing)}")

        rows: list[dict[str, str]] = []
        for row_idx in range(min_row + 1, max_row + 1):
            item = {
                "reviewee": self._to_string(worksheet.cell(row=row_idx, column=header_map["reviewee"]).value),
                "email": self._to_string(worksheet.cell(row=row_idx, column=header_map["email"]).value),
                "subco": self._to_string(worksheet.cell(row=row_idx, column=header_map["subco"]).value),
                "horasQuincena1": self._to_string(
                    worksheet.cell(row=row_idx, column=header_map["horasquincena1"]).value
                ),
                "horasQuincena2": self._to_string(
                    worksheet.cell(row=row_idx, column=header_map["horasquincena2"]).value
                ),
            }
            if all(value == "" for value in item.values()):
                continue
            rows.append(item)

        return rows

    def _build_header_map(self, headers: list[object], min_col: int) -> dict[str, int]:
        mapping: dict[str, int] = {}
        for offset, header in enumerate(headers):
            if header is None:
                continue
            normalized = str(header).strip().casefold()
            if normalized and normalized not in mapping:
                mapping[normalized] = offset

        return {
            header_key: min_col + mapping[header_key]
            for header_key in REQUIRED_HEADER_KEYS
            if header_key in mapping
        }

    def _to_string(self, value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            return str(int(value)) if value.is_integer() else f"{value:g}"
        return str(value).strip()
