from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
import json
import re

from openpyxl import load_workbook
from openpyxl.styles import Color

SHEET_NAME = "Timecards Aprobadas"
HEADER_ROW_TIME_PERIOD = 3
REVIEWEE_COLUMN = 1

APPROVED_FONT_RGB = "FF00B050"

VALIDATION_OK = "OK"
VALIDATION_FAILED = "Failed"
VALIDATION_NOT_FOUND = "NotFound"

ACTION_APPROVE = "Approve"
ACTION_REJECT = "Reject"
ACTION_NA = "NA"

MESSAGE_REVIEWEE_NOT_FOUND = "Reviewee not Found."
MESSAGE_DUPLICATE = "Timecard enviada por duplicado."
MESSAGE_AMOUNT_NOT_AVAILABLE = "totalEstimatedAmount no disponible en Excel. Revisar cálculo en el fichero."
MESSAGE_INCORRECT_TIMECARD = "Timecard incorrecta. Revisar y Rechzarla"
MESSAGE_INVALID_INPUT = "Input item missing required fields or empty values."

REQUIRED_INPUT_FIELDS = ("beelineRevieweeName", "timePeriod", "totalEstimatedAmount")


@dataclass(frozen=True)
class BeelineValidationResult:
    revieweeBeelineName: str
    timePeriod: str
    validationStatus: str
    message: str
    beelineAction: str


@dataclass(frozen=True)
class BeelineValidationRunResult:
    workbook_path: Path
    input_json_path: Path
    output_path: Path
    validated_rows: int
    approved_rows: int
    rejected_rows: int
    not_found_rows: int


class BeelineTimesheetValidationService:
    """Validate Beeline timecards against the control workbook and export a JSON decision list."""

    def run(
        self,
        workbook_path: Path,
        input_json_path: Path,
        output_file_name: str,
        output_root: Path,
    ) -> BeelineValidationRunResult:
        normalized_output_name = self._normalize_output_name(output_file_name)
        input_items = self._read_input_json(input_json_path)

        workbook = load_workbook(workbook_path)
        workbook_data_only = load_workbook(workbook_path, data_only=True)
        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise ValueError(f"Worksheet not found in workbook: {SHEET_NAME}")
            if SHEET_NAME not in workbook_data_only.sheetnames:
                raise ValueError(f"Worksheet not found in workbook: {SHEET_NAME}")
            worksheet = workbook[SHEET_NAME]
            worksheet_data_only = workbook_data_only[SHEET_NAME]

            validations: list[BeelineValidationResult] = []
            for input_item in input_items:
                result = self._validate_item(input_item, worksheet, worksheet_data_only)
                validations.append(result)

            workbook.save(workbook_path)
        finally:
            workbook_data_only.close()
            workbook.close()

        output_root.mkdir(parents=True, exist_ok=True)
        output_path = output_root / f"{normalized_output_name}.json"
        output_payload = [
            {
                "revieweeBeelineName": row.revieweeBeelineName,
                "timePeriod": row.timePeriod,
                "validationStatus": row.validationStatus,
                "message": row.message,
                "beelineAction": row.beelineAction,
            }
            for row in validations
        ]
        output_path.write_text(json.dumps(output_payload, ensure_ascii=False, indent=2), encoding="utf-8")

        approved_rows = sum(1 for row in validations if row.validationStatus == VALIDATION_OK)
        rejected_rows = sum(
            1
            for row in validations
            if row.validationStatus == VALIDATION_FAILED and row.beelineAction == ACTION_REJECT
        )
        not_found_rows = sum(1 for row in validations if row.validationStatus == VALIDATION_NOT_FOUND)

        return BeelineValidationRunResult(
            workbook_path=workbook_path,
            input_json_path=input_json_path,
            output_path=output_path,
            validated_rows=len(validations),
            approved_rows=approved_rows,
            rejected_rows=rejected_rows,
            not_found_rows=not_found_rows,
        )

    def _read_input_json(self, input_json_path: Path) -> list[dict[str, object]]:
        raw = json.loads(input_json_path.read_text(encoding="utf-8"))
        if not isinstance(raw, list):
            raise ValueError("input JSON must be an array of objects")

        normalized: list[dict[str, object]] = []
        for item in raw:
            if isinstance(item, dict):
                normalized.append(item)
            else:
                normalized.append({})

        return normalized

    def _validate_item(self, item: dict[str, object], worksheet, worksheet_data_only) -> BeelineValidationResult:
        reviewee = self._as_string(item.get("beelineRevieweeName"))
        time_period = self._as_string(item.get("timePeriod"))
        total_estimated_amount = self._as_string(item.get("totalEstimatedAmount"))

        if not self._has_all_required_fields(item):
            return BeelineValidationResult(
                revieweeBeelineName=reviewee,
                timePeriod=time_period,
                validationStatus=VALIDATION_FAILED,
                message=MESSAGE_INVALID_INPUT,
                beelineAction=ACTION_REJECT,
            )

        row_idx = self._find_reviewee_row(worksheet, reviewee)
        if row_idx is None:
            return BeelineValidationResult(
                revieweeBeelineName=reviewee,
                timePeriod=time_period,
                validationStatus=VALIDATION_NOT_FOUND,
                message=MESSAGE_REVIEWEE_NOT_FOUND,
                beelineAction=ACTION_NA,
            )

        col_idx = self._find_time_period_column(worksheet, worksheet_data_only, time_period)
        if col_idx is None:
            return BeelineValidationResult(
                revieweeBeelineName=reviewee,
                timePeriod=time_period,
                validationStatus=VALIDATION_NOT_FOUND,
                message=MESSAGE_AMOUNT_NOT_AVAILABLE,
                beelineAction=ACTION_NA,
            )

        target_cell = worksheet.cell(row=row_idx, column=col_idx)
        if self._is_approved_style(target_cell):
            return BeelineValidationResult(
                revieweeBeelineName=reviewee,
                timePeriod=time_period,
                validationStatus=VALIDATION_FAILED,
                message=MESSAGE_DUPLICATE,
                beelineAction=ACTION_REJECT,
            )

        excel_amount = self._as_string(target_cell.value)
        if excel_amount == "":
            return BeelineValidationResult(
                revieweeBeelineName=reviewee,
                timePeriod=time_period,
                validationStatus=VALIDATION_NOT_FOUND,
                message=MESSAGE_AMOUNT_NOT_AVAILABLE,
                beelineAction=ACTION_NA,
            )

        excel_numeric_amount = self._parse_amount(excel_amount)
        input_numeric_amount = self._parse_amount(total_estimated_amount)

        if excel_numeric_amount is not None and input_numeric_amount is not None and excel_numeric_amount == input_numeric_amount:
            self._mark_approved_style(target_cell)
            return BeelineValidationResult(
                revieweeBeelineName=reviewee,
                timePeriod=time_period,
                validationStatus=VALIDATION_OK,
                message="",
                beelineAction=ACTION_APPROVE,
            )

        return BeelineValidationResult(
            revieweeBeelineName=reviewee,
            timePeriod=time_period,
            validationStatus=VALIDATION_FAILED,
            message=MESSAGE_INCORRECT_TIMECARD,
            beelineAction=ACTION_REJECT,
        )

    def _normalize_output_name(self, output_file_name: str) -> str:
        name = output_file_name.strip()
        if name.lower().endswith(".json"):
            name = name[:-5]
        if not name:
            raise ValueError("output-file-name must not be empty")
        if any(char in name for char in "\\/:*?\"<>|"):
            raise ValueError("output-file-name contains invalid filesystem characters")
        return name

    def _has_all_required_fields(self, item: dict[str, object]) -> bool:
        for key in REQUIRED_INPUT_FIELDS:
            if key not in item:
                return False
            value = item[key]
            if not isinstance(value, str):
                return False
            if value == "":
                return False
        return True

    def _find_reviewee_row(self, worksheet, reviewee_name: str) -> int | None:
        for row_idx in range(1, worksheet.max_row + 1):
            cell_value = self._as_string(worksheet.cell(row=row_idx, column=REVIEWEE_COLUMN).value)
            if cell_value == reviewee_name:
                return row_idx
        return None

    def _find_time_period_column(self, worksheet, worksheet_data_only, time_period: str) -> int | None:
        target_normalized = self._normalize_time_period_string(time_period)
        max_col = max(worksheet.max_column, worksheet_data_only.max_column)
        resolved_formula_periods: dict[int, str] = {}

        for col_idx in range(1, max_col + 1):
            raw_header = worksheet.cell(row=HEADER_ROW_TIME_PERIOD, column=col_idx).value
            data_only_header = worksheet_data_only.cell(row=HEADER_ROW_TIME_PERIOD, column=col_idx).value
            candidates = self._collect_header_candidates(
                raw_header=raw_header,
                data_only_header=data_only_header,
                col_idx=col_idx,
                worksheet=worksheet,
                worksheet_data_only=worksheet_data_only,
                resolved_formula_periods=resolved_formula_periods,
            )

            for candidate in candidates:
                if candidate == time_period:
                    return col_idx

                candidate_normalized = self._normalize_time_period_string(candidate)
                if target_normalized is not None and candidate_normalized == target_normalized:
                    return col_idx
        return None

    def _collect_header_candidates(
        self,
        raw_header: object,
        data_only_header: object,
        col_idx: int,
        worksheet,
        worksheet_data_only,
        resolved_formula_periods: dict[int, str],
    ) -> list[str]:
        candidates: list[str] = []

        for header_value in (raw_header, data_only_header):
            if isinstance(header_value, str):
                text = header_value.strip()
                if text and not text.startswith("="):
                    candidates.append(text)

        if isinstance(raw_header, str) and raw_header.strip().startswith("="):
            derived = self._resolve_formula_period_from_previous_column(
                col_idx=col_idx,
                worksheet=worksheet,
                worksheet_data_only=worksheet_data_only,
                resolved_formula_periods=resolved_formula_periods,
            )
            if derived:
                candidates.append(derived)

        return candidates

    def _resolve_formula_period_from_previous_column(
        self,
        col_idx: int,
        worksheet,
        worksheet_data_only,
        resolved_formula_periods: dict[int, str],
    ) -> str | None:
        if col_idx in resolved_formula_periods:
            return resolved_formula_periods[col_idx]

        if col_idx <= 1:
            return None

        previous_candidates = self._collect_header_candidates(
            raw_header=worksheet.cell(row=HEADER_ROW_TIME_PERIOD, column=col_idx - 1).value,
            data_only_header=worksheet_data_only.cell(row=HEADER_ROW_TIME_PERIOD, column=col_idx - 1).value,
            col_idx=col_idx - 1,
            worksheet=worksheet,
            worksheet_data_only=worksheet_data_only,
            resolved_formula_periods=resolved_formula_periods,
        )

        for previous_candidate in previous_candidates:
            parsed = self._parse_time_period_to_dates(previous_candidate)
            if parsed is None:
                continue

            start_date, end_date = parsed
            derived_start = start_date + timedelta(days=7)
            derived_end = end_date + timedelta(days=7)
            derived = self._format_date_range(derived_start, derived_end)
            resolved_formula_periods[col_idx] = derived
            return derived

        return None

    def _as_string(self, value: object) -> str:
        if value is None:
            return ""
        return str(value)

    def _normalize_time_period_string(self, value: str) -> str | None:
        parsed = self._parse_time_period_to_dates(value)
        if parsed is None:
            return None
        start_date, end_date = parsed
        return self._format_date_range(start_date, end_date)

    def _parse_time_period_to_dates(self, value: str) -> tuple[date, date] | None:
        normalized = value.strip().replace("–", "-").replace("—", "-")
        parts = [part.strip() for part in normalized.split("-")]
        if len(parts) != 2:
            return None

        start_date = self._parse_us_date(parts[0])
        end_date = self._parse_us_date(parts[1])
        if start_date is None or end_date is None:
            return None

        return start_date, end_date

    def _parse_us_date(self, value: str) -> date | None:
        match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", value)
        if not match:
            return None

        month = int(match.group(1))
        day = int(match.group(2))
        year = int(match.group(3))
        try:
            return date(year, month, day)
        except ValueError:
            return None

    def _format_date_range(self, start_date: date, end_date: date) -> str:
        return (
            f"{start_date.month}/{start_date.day}/{start_date.year}"
            f"-{end_date.month}/{end_date.day}/{end_date.year}"
        )

    def _parse_amount(self, value: str) -> Decimal | None:
        text = value.strip()
        if text == "":
            return None

        # Keep only numeric content and common separators to support values like EUR strings.
        cleaned = re.sub(r"[^0-9,.-]", "", text)
        if cleaned in {"", ".", ",", "-", "-.", "-,"}:
            return None

        normalized = self._normalize_numeric_separators(cleaned)
        try:
            return Decimal(normalized)
        except InvalidOperation:
            return None

    def _normalize_numeric_separators(self, value: str) -> str:
        last_dot = value.rfind(".")
        last_comma = value.rfind(",")

        if last_dot != -1 and last_comma != -1:
            # If comma appears after dot, assume EU style (1.234,56). Otherwise US style (1,234.56).
            if last_comma > last_dot:
                return value.replace(".", "").replace(",", ".")
            return value.replace(",", "")

        if last_comma != -1:
            return value.replace(",", ".")

        return value

    def _is_approved_style(self, cell) -> bool:
        font = cell.font
        if font is None or not bool(font.bold):
            return False
        color = font.color
        if color is None:
            return False
        return color.type == "rgb" and color.rgb == APPROVED_FONT_RGB

    def _mark_approved_style(self, cell) -> None:
        updated_font = copy(cell.font)
        updated_font.bold = True
        updated_font.color = Color(rgb=APPROVED_FONT_RGB)
        cell.font = updated_font
